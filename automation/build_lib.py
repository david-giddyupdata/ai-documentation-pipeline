"""Shared Word-editing logic: rewrite one exercise's instruction table in the draft.

Step text mini-markup: **bold** for UI labels, *italic*, {field} for field names (Code In Text style).
Each step is a dict:
    items   optional ordered list of ("line"|"bullet"|"keep", value); overrides the three below
    lines   numbered paragraphs (usually one)
    bullets sub-bullets (List Bullet 2)
    keep    paragraph elements copied verbatim from the old table (e.g. formula lines)
    shot    annotated screenshot stem
"""
import copy
import re
from pathlib import Path

import docx
from docx.oxml.ns import qn
from docx.shared import Inches
from PIL import Image

HERE_SHOTS = Path(__file__).resolve().parent / "screenshots_annotated"
MAX_W = 5.5   # inches; right column is 5.9"
SCALE = 0.8   # inches per 96 css px; matches the size of the hand-made screenshots
DPR = 2       # screenshots captured at devicePixelRatio 2
CODE_STYLE = "Code In Text [PACKT]"   # style name inherited from the master document's Word template;
                                      # renaming it would mean changing the master


def text_of(el):
    return "".join(t.text or "" for t in el.iter(qn("w:t"))).strip()


def set_runs(p_el, parent, text):
    """Replace a paragraph's runs with marked-up text, keeping its pPr."""
    for r in list(p_el):
        if r.tag != qn("w:pPr"):
            p_el.remove(r)
    para = docx.text.paragraph.Paragraph(p_el, parent)
    for tok in re.split(r"(\*\*.+?\*\*|\*[^*]+\*|\{[^}]+\})", text):
        if not tok:
            continue
        if tok.startswith("**"):
            para.add_run(tok[2:-2]).bold = True
        elif tok.startswith("*"):
            para.add_run(tok[1:-1]).italic = True
        elif tok.startswith("{"):
            para.add_run(tok, style=CODE_STYLE)
        else:
            para.add_run(tok)


def picture_width(path, scale=SCALE, max_w=MAX_W):
    im = Image.open(path)
    return min(im.width / DPR / 96 * scale, max_w)


def find_table(d, heading):
    body = list(d.element.body.iterchildren())
    h = next(i for i, el in enumerate(body) if el.tag == qn("w:p") and text_of(el) == heading)
    tbl_el = next(el for el in body[h:] if el.tag == qn("w:tbl"))
    return body, tbl_el, docx.table.Table(tbl_el, d._body)


def rebuild(d, heading, steps, shots, completed, shot_scale=None, completed_w=7.5,
            first_starts="Click the + sign"):
    """Replace every row but the last (Completed Results) of the heading's first table.
    Returns the old rows (so callers can pull `keep` paragraphs from them)."""
    shot_scale = shot_scale or {}
    body, tbl_el, t = find_table(d, heading)
    rows = list(t.rows)
    assert text_of(rows[0]._tr).startswith(first_starts), f"unexpected table under {heading}"
    completed_tr = rows[-1]._tr
    if "Completed Results" in text_of(completed_tr):
        step_rows, completed_tbl = rows[:-1], None
    else:  # Completed Results is a separate one-cell table right after (e.g. Sparklines)
        step_rows = rows
        completed_tbl = next(el for el in body[body.index(tbl_el) + 1:] if el.tag == qn("w:tbl"))
        assert "Completed Results" in text_of(completed_tbl), f"no Completed Results for {heading}"

    # templates: the first step paragraph carries the exercise's own list (numId with a
    # restart override); later steps in the master used a shared numId that keeps counting
    # across exercises, so every step is cloned from the first.
    first_p = rows[0].cells[0].paragraphs[0]._p
    step_p = copy.deepcopy(first_p)
    bullet_p = next((p._p for r in rows for p in r.cells[0].paragraphs if p.style.name == "List Bullet 2"), None)
    if bullet_p is None:  # table has no sub-bullets: borrow the first one in any table
        bullet_p = next(p for p in d.element.body.iter(qn("w:p"))
                        if (s := p.find(qn("w:pPr") + "/" + qn("w:pStyle"))) is not None
                        and s.get(qn("w:val")) == "ListBullet2" and p.getparent().tag == qn("w:tc"))
    bullet_p = copy.deepcopy(bullet_p)
    row_tpl = copy.deepcopy(rows[0]._tr)
    old_rows = [copy.deepcopy(r._tr) for r in step_rows]

    for r in step_rows:
        tbl_el.remove(r._tr)

    for st in steps:
        tr = copy.deepcopy(row_tpl)
        if completed_tbl is None:
            completed_tr.addprevious(tr)
        else:
            tbl_el.append(tr)
        left, right = tr.findall(qn("w:tc"))
        lcell = docx.table._Cell(left, t)
        for el in list(left):
            if el.tag == qn("w:p"):
                left.remove(el)
        items = st.get("items") or ([("line", x) for x in st.get("lines", [])] +
                                    [("keep", x) for x in st.get("keep", [])] +
                                    [("bullet", x) for x in st.get("bullets", [])])
        for kind, val in items:
            if kind == "keep":
                left.append(copy.deepcopy(val))
                continue
            p = copy.deepcopy(step_p if kind == "line" else bullet_p)
            left.append(p)
            set_runs(p, lcell, val)
        rps = right.findall(qn("w:p"))
        for el in rps[1:]:
            right.remove(el)
        for r in rps[0].findall(qn("w:r")):
            rps[0].remove(r)
        if st.get("shot"):
            path = shots / f"{st['shot']}.png"
            cell = docx.table._Cell(right, t)
            cell.paragraphs[0].add_run().add_picture(
                str(path), width=Inches(picture_width(path, shot_scale.get(st["shot"], SCALE))))

    if completed_tbl is not None:
        completed_tr = completed_tbl.find(qn("w:tr"))
    cell = docx.table._Cell(completed_tr.findall(qn("w:tc"))[0], t)
    for p in cell.paragraphs:
        for r in p._p.findall(qn("w:r")):
            if r.find(".//" + qn("w:drawing")) is not None:
                p._p.remove(r)
    path = shots / f"{completed}.png"
    cell.paragraphs[-1].add_run().add_picture(str(path), width=Inches(picture_width(path, 1.0, completed_w)))
    return old_rows, tbl_el


# Step numbering. Main steps in the master use abstractNum 6 (many numIds, each restarting at 1)
# or numId 56 / abstractNum 11, which never restarts, so counts ran on across exercises
# (1, 48, 49...). Sub-lists use other abstracts (mostly bullets) and are left alone.
MAIN_ABS = ("6", "11")
STEP_ABS = "6"


def _abstracts(d):
    numbering = d.part.numbering_part.element
    return numbering, {n.get(qn("w:numId")): n.find(qn("w:abstractNumId")).get(qn("w:val"))
                       for n in numbering.findall(qn("w:num"))}


def _num_id(p):
    n = p.find(qn("w:pPr") + "/" + qn("w:numPr") + "/" + qn("w:numId"))
    return n.get(qn("w:val")) if n is not None else None


def _ilvl(p):
    n = p.find(qn("w:pPr") + "/" + qn("w:numPr") + "/" + qn("w:ilvl"))
    return n.get(qn("w:val")) if n is not None else "0"


def step_groups(d):
    """[(heading chain, [main-step paragraphs])] for every numbered list, in tables or not.
    Lists under the same heading chain (e.g. Connect to S3 and Union's two tables) form one
    group so their numbering continues; a Commentary list under its own Heading 4 doesn't."""
    _, abs_of = _abstracts(d)
    is_step = lambda p: abs_of.get(_num_id(p)) in MAIN_ABS and _ilvl(p) == "0"
    heads, groups = {}, {}
    for el in d.element.body.iterchildren():
        if el.tag == qn("w:p"):
            sty = el.find(qn("w:pPr") + "/" + qn("w:pStyle"))
            sid = sty.get(qn("w:val")) if sty is not None else ""
            if sid.startswith("Heading") and sid[7:].isdigit() and text_of(el):
                lvl = int(sid[7:])
                heads = {k: v for k, v in heads.items() if k < lvl}
                heads[lvl] = text_of(el)
            steps = [el] if is_step(el) else []
        elif el.tag == qn("w:tbl"):
            steps = [p for p in el.iter(qn("w:p")) if is_step(p)]
        else:
            steps = []
        if steps:
            groups.setdefault(" > ".join(heads[k] for k in sorted(heads)), []).extend(steps)
    return list(groups.items())


def fix_numbering(d):
    """Give each step group its own list (abstractNum 6, restart at 1) so every exercise reads 1, 2, 3..."""
    numbering, abs_of = _abstracts(d)
    nums = numbering.findall(qn("w:num"))
    next_id = max(int(n) for n in abs_of) + 1
    for _, paras in step_groups(d):
        num = numbering.makeelement(qn("w:num"), {qn("w:numId"): str(next_id)})
        a = num.makeelement(qn("w:abstractNumId"), {qn("w:val"): STEP_ABS})
        lo = num.makeelement(qn("w:lvlOverride"), {qn("w:ilvl"): "0"})
        lo.append(lo.makeelement(qn("w:startOverride"), {qn("w:val"): "1"}))
        num.extend([a, lo])
        nums[-1].addnext(num)
        nums[-1] = num
        for p in paras:
            p.find(qn("w:pPr") + "/" + qn("w:numPr") + "/" + qn("w:numId")).set(qn("w:val"), str(next_id))
        next_id += 1


def replace_after(d, tbl_el, mapping, stop_heading_styles=("Heading2",)):
    """Replace text in single runs between a table and the next Heading 2."""
    body = list(d.element.body.iterchildren())
    found = set()
    for el in body[body.index(tbl_el) + 1:]:
        if el.tag == qn("w:p"):
            sty = el.find(".//" + qn("w:pStyle"))
            if sty is not None and sty.get(qn("w:val")) in stop_heading_styles and text_of(el):
                break
        for tnode in el.iter(qn("w:t")):
            for old, new in mapping.items():
                if tnode.text and old in tnode.text:
                    tnode.text = tnode.text.replace(old, new)
                    found.add(old)
    missing = set(mapping) - found
    if missing:
        print("WARNING: not found in single runs:", missing)


# Surgical edits for hand-written exercises (A, B, C): change single paragraphs, keep the rest.
def exercise_paras(d, heading):
    """Every paragraph (tables included) from the heading to the next Heading 1/2."""
    body = list(d.element.body.iterchildren())
    h = next(i for i, el in enumerate(body) if el.tag == qn("w:p") and text_of(el) == heading)
    out = []
    for el in body[h + 1:]:
        sty = el.find(qn("w:pPr") + "/" + qn("w:pStyle")) if el.tag == qn("w:p") else None
        if sty is not None and sty.get(qn("w:val")) in ("Heading1", "Heading2") and text_of(el):
            break
        out.extend(el.iter(qn("w:p")))
    return out


def find_para(d, heading, starts):
    hits = [p for p in exercise_paras(d, heading) if text_of(p).startswith(starts)]
    assert len(hits) == 1, f"{heading}: {len(hits)} paragraphs start with {starts!r}"
    return hits[0]


def set_para(d, heading, starts, markup):
    p = find_para(d, heading, starts)
    set_runs(p, d._body, markup)
    return p


def add_bullet_after(d, heading, starts, markup):
    """Insert a List Bullet 2 paragraph (cloned from the exercise) after the matching paragraph."""
    return add_bullet_after_el(d, heading, find_para(d, heading, starts), markup)


def delete_para(d, heading, starts):
    p = find_para(d, heading, starts)
    p.getparent().remove(p)


def make_bullet(d, heading, starts):
    """Turn a numbered paragraph into a List Bullet 2 sub-bullet (keeps its runs)."""
    p = find_para(d, heading, starts)
    tpl = next(b for b in exercise_paras(d, heading)
               if (s := b.find(qn("w:pPr") + "/" + qn("w:pStyle"))) is not None
               and s.get(qn("w:val")) == "ListBullet2")
    p.replace(p.find(qn("w:pPr")), copy.deepcopy(tpl.find(qn("w:pPr"))))
    return p


def add_bullets_after(d, heading, starts, markups):
    p = find_para(d, heading, starts)
    for m in markups:
        p = add_bullet_after_el(d, heading, p, m)


def add_bullet_after_el(d, heading, p, markup):
    tpl = next(b for b in exercise_paras(d, heading)
               if (s := b.find(qn("w:pPr") + "/" + qn("w:pStyle"))) is not None
               and s.get(qn("w:val")) == "ListBullet2")
    b = copy.deepcopy(tpl)
    p.addnext(b)
    set_runs(b, d._body, markup)
    return b



def _step_tpl(d, heading):
    return next(p for p in exercise_paras(d, heading)
                if (s := p.find(qn("w:pPr") + "/" + qn("w:pStyle"))) is not None
                and s.get(qn("w:val")) == "ListNumber" and p.getparent().tag == qn("w:tc"))


def make_step(d, heading, starts):
    """Turn a bullet into a numbered main step (keeps its runs)."""
    p = find_para(d, heading, starts)
    p.replace(p.find(qn("w:pPr")), copy.deepcopy(_step_tpl(d, heading).find(qn("w:pPr"))))
    return p


def add_step_after(d, heading, starts, markup):
    p = find_para(d, heading, starts)
    s = copy.deepcopy(_step_tpl(d, heading))
    p.addnext(s)
    set_runs(s, d._body, markup)
    return s


def row_image_part(d, heading, row_starts):
    """The image part of the screenshot in the instruction row whose text starts with row_starts."""
    _, tbl_el, _ = find_table(d, heading)
    tr = next(r for r in tbl_el.findall(qn("w:tr")) if text_of(r).startswith(row_starts))
    blip = next(tr.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}blip"))
    return d.part.related_parts[blip.get(qn("r:embed"))]


def add_box_to_image(part, rect, pad=4, width=3, radius=6):
    """Draw a David-style #F4AC08 rounded box (soft shadow) on an embedded image, in place."""
    import io
    from PIL import ImageDraw, ImageFilter
    im = Image.open(io.BytesIO(part.blob)).convert("RGBA")
    x0, y0, x1, y1 = rect[0] - pad, rect[1] - pad, rect[2] + pad, rect[3] + pad
    sh = Image.new("RGBA", im.size, (0, 0, 0, 0))
    ImageDraw.Draw(sh).rounded_rectangle((x0 + 2, y0 + 2, x1 + 2, y1 + 2), radius, outline=(0, 0, 0, 90), width=width)
    im = Image.alpha_composite(im, sh.filter(ImageFilter.GaussianBlur(2)))
    ImageDraw.Draw(im).rounded_rectangle((x0, y0, x1, y1), radius, outline=(0xF4, 0xAC, 0x08, 255), width=width)
    buf = io.BytesIO()
    im.convert("RGB").save(buf, "PNG")
    part._blob = buf.getvalue()


def media_part(d, name):
    """An embedded image part by its media file name (e.g. 'image6.png')."""
    return next(p for p in d.part.package.parts if str(p.partname).endswith("/media/" + name))


def swap_image(d, name, png, width_in=None):
    """Replace an embedded image (media file name, e.g. 'image22.png') with a new PNG everywhere it
    is used. Each use keeps its width (or takes width_in) and gets the new aspect ratio. Session 6
    retakes for the hand-written sections (A, B, C) use this; the media part is shared, so every
    paragraph that shows the old picture shows the new one."""
    part = media_part(d, name)
    assert str(part.partname).lower().endswith(".png"), part.partname
    part._blob = Path(png).read_bytes()
    w, h = Image.open(png).size
    A_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    WP = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}"
    rids = {rid for rid, rel in d.part.rels.items() if not rel.is_external and rel.target_part is part}
    n = 0
    for blip in d.element.body.iter(A_NS + "blip"):
        if blip.get(qn("r:embed")) not in rids:
            continue
        drawing = blip
        while drawing is not None and drawing.tag != qn("w:drawing"):
            drawing = drawing.getparent()
        ext = next(drawing.iter(WP + "extent"))
        cx = int(width_in * 914400) if width_in else int(ext.get("cx"))
        cy = int(cx * h / w)
        ext.set("cx", str(cx)); ext.set("cy", str(cy))
        for e in drawing.iter(A_NS + "ext"):
            if e.get("cx") is not None:
                e.set("cx", str(cx)); e.set("cy", str(cy))
        n += 1
    assert n, f"{name}: no uses found"
    return n


def blur_schema(part, prefix="quicksight_", pad=2):
    """Blur whatever follows `prefix` in every OCR'd word that starts with it (David: the schema
    name may change, so screenshots never show it). Returns the boxes blurred, in image px."""
    import io
    import pytesseract
    from PIL import ImageFilter
    import os
    pytesseract.pytesseract.tesseract_cmd = os.environ.get(
        "TESSERACT_CMD", r"C:\Program Files\Tesseract-OCR\tesseract.exe")
    im = Image.open(io.BytesIO(part.blob)).convert("RGB")
    k = 3
    big = im.resize((im.width * k, im.height * k))
    data = pytesseract.image_to_data(big, output_type=pytesseract.Output.DICT)
    boxes = []
    for i, w in enumerate(data["text"]):
        if not w.lower().startswith(prefix) or len(w) <= len(prefix):
            continue
        x, y, ww, hh = (data[c][i] // k for c in ("left", "top", "width", "height"))
        xs = x + round(ww * len(prefix) / len(w)) - 2  # small fonts: start a little early
        boxes.append((xs, y - pad, x + ww + pad + 2, y + hh + pad))
    for b in boxes:
        region = im.crop(b).filter(ImageFilter.GaussianBlur(4))
        im.paste(region, b[:2])
    buf = io.BytesIO()
    im.save(buf, "PNG")
    part._blob = buf.getvalue()
    return boxes


def swap_run(d, heading, starts, old, markup):
    """Replace one run (exact text `old`) inside a paragraph with marked-up runs, keeping the
    paragraph's other runs (line breaks, formulas) untouched. Plain parts copy the old run's rPr."""
    p = find_para(d, heading, starts)
    run = next(r for r in p.findall(qn("w:r")) if text_of(r) == old.strip() and (r.find(qn("w:t")).text == old))
    rpr = run.find(qn("w:rPr"))
    anchor = run
    for tok in re.split(r"(\*\*.+?\*\*)", markup):
        if not tok:
            continue
        r = run.makeelement(qn("w:r"), {})
        if rpr is not None:
            r.append(copy.deepcopy(rpr))
        if tok.startswith("**"):
            tok = tok[2:-2]
            rp = r.find(qn("w:rPr"))
            if rp is None:
                rp = r.makeelement(qn("w:rPr"), {})
                r.insert(0, rp)
            if rp.find(qn("w:b")) is None:
                rp.append(rp.makeelement(qn("w:b"), {}))
        t = r.makeelement(qn("w:t"), {})
        t.text = tok
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        r.append(t)
        anchor.addnext(r)
        anchor = r
    p.remove(run)


# ---- Layout passes (David's review rules, 2026-10-13; steering/style.md "David's Review Rules") ----
# Run last in make_draft.py, after every exercise and fix_numbering.
EMU = 914400
CELL_MAX = 5.57   # inches: widest picture in a two-column instruction row
ROW_MAX = 9.86    # inches: widest picture in a full-width row (Completed Results) or in body text
MAX_H = 6.6       # inches: tallest picture that still fits on a page with its row
_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_WP = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}"


def _set_size(drawing, cx, cy):
    for e in drawing.iter(_WP + "extent"):
        e.set("cx", str(cx)); e.set("cy", str(cy))
    for e in drawing.iter(_A + "ext"):
        if e.get("cx") is not None:
            e.set("cx", str(cx)); e.set("cy", str(cy))


def fit_images(d, min_scale=0.895):   # Word shows 89.5% as 90%
    """Size every picture we captured (saved at 192 dpi, so 100% = on-screen size) at 100%, capped at
    CELL_MAX in a two-column row, ROW_MAX elsewhere, and MAX_H tall. David's own 96-dpi images keep
    their size. Returns the pictures that can't reach min_scale; they need cropping, a cut-out, a tear
    or a recapture rather than shrinking."""
    import io
    short, chain = [], {}
    for el in d.element.body.iterchildren():
        if el.tag == qn("w:p"):
            st = el.find(qn("w:pPr") + "/" + qn("w:pStyle"))
            m = re.fullmatch(r"Heading([1-3])", st.get(qn("w:val"))) if st is not None else None
            if m and text_of(el):
                lvl = int(m.group(1))
                chain = {k: v for k, v in chain.items() if k < lvl}
                chain[lvl] = text_of(el)
        for drawing in el.iter(qn("w:drawing")):
            blip = next(drawing.iter(_A + "blip"), None)
            if blip is None:
                continue
            im = Image.open(io.BytesIO(d.part.related_parts[blip.get(qn("r:embed"))].blob))
            dpi = im.info.get("dpi", (96, 96))[0]
            if dpi < 150:
                continue
            tr = next(drawing.iterancestors(qn("w:tr")), None)
            max_w = CELL_MAX if tr is not None and len(tr.findall(qn("w:tc"))) > 1 else ROW_MAX
            nat = im.width / dpi
            by_h = MAX_H * im.width / im.height
            target = min(nat, max_w, by_h)
            _set_size(drawing, int(target * EMU), int(target * EMU * im.height / im.width))
            if target < nat * min_scale - 1e-6:
                where = chain.get(3) if chain.get(1, "").startswith("Addendum") else chain.get(2, chain.get(1, ""))
                short.append(dict(exercise=where, row=(text_of(tr) if tr is not None else "")[:80],
                                  scale=round(target / nat * 100), width=round(target, 2),
                                  limit="height" if target == by_h else "width"))
    return short


def space_after_pictures(d, pt=6):
    """A little space between a picture and the cell border below it: space after every picture
    paragraph that ends a table cell (David: images touched the bottom border)."""
    from docx.shared import Pt
    n = 0
    for tc in d.element.body.iter(qn("w:tc")):
        ps = tc.findall(qn("w:p"))
        if ps and ps[-1].find(".//" + qn("w:drawing")) is not None:
            docx.text.paragraph.Paragraph(ps[-1], d._body).paragraph_format.space_after = Pt(pt)
            n += 1
    return n


def page_breaks_to_headings(d):
    """A paragraph holding only a page break right before a heading spills onto its own page when the
    page above is full, leaving a blank page (p141 in session 7). Put the break on the heading instead
    (Page break before) and drop that paragraph and any empty paragraphs just above it."""
    body = d.element.body

    def empty(p):
        return (p is not None and p.tag == qn("w:p") and not text_of(p) and p.find(".//" + qn("w:drawing")) is None
                and p.find(qn("w:pPr") + "/" + qn("w:sectPr")) is None)

    n = 0
    for p in list(body.iterchildren(qn("w:p"))):
        brs = list(p.iter(qn("w:br")))
        if not brs or not empty(p) or any(b.get(qn("w:type")) != "page" for b in brs):
            continue
        nxt = p.getnext()
        st = nxt.find(qn("w:pPr") + "/" + qn("w:pStyle")) if nxt is not None and nxt.tag == qn("w:p") else None
        if st is None or not re.fullmatch(r"Heading[1-9]", st.get(qn("w:val"))):
            continue
        prev = p.getprevious()
        while empty(prev) and not list(prev.iter(qn("w:br"))):
            above = prev.getprevious()
            body.remove(prev)
            prev = above
        body.remove(p)
        docx.text.paragraph.Paragraph(nxt, d._body).paragraph_format.page_break_before = True
        n += 1
    return n


def keep_picture_rows_together(d, max_h=MAX_H):
    """Table rows with a picture don't break across pages, so "Completed Results:" stays with its image
    and a step stays with its screenshot. Rows with a picture taller than max_h are skipped (Word cuts
    off an unbreakable row that is taller than the page)."""
    later = [qn(t) for t in ("w:trHeight", "w:tblHeader", "w:tblCellSpacing", "w:jc", "w:hidden",
                             "w:ins", "w:del", "w:trPrChange")]
    n = 0
    for tr in d.element.body.iter(qn("w:tr")):
        hs = [int(e.get("cy")) / EMU for e in tr.iter(_WP + "extent")]
        if not hs or max(hs) > max_h:
            continue
        trPr = tr.get_or_add_trPr()
        if trPr.find(qn("w:cantSplit")) is None:
            cs = trPr.makeelement(qn("w:cantSplit"), {})
            nxt = next((c for c in trPr if c.tag in later), None)
            nxt.addprevious(cs) if nxt is not None else trPr.append(cs)
            n += 1
    return n


def merge_completed_results(d):
    """A "Completed Results:" row whose label is its own paragraph (e.g. the master's Pivot Table, a Heading 4)
    becomes the standard form: one centered paragraph, label (Heading4Char) + line break + picture. The separate
    heading paragraph made the Pivot row taller than a page, so Word split it: label on one page, picture on
    the next (student walkthrough, session 9)."""
    n = 0
    for tr in d.element.body.iter(qn("w:tr")):
        tc = tr.find(qn("w:tc"))
        ps = tc.findall(qn("w:p")) if tc is not None else []
        if len(ps) != 2:
            continue
        label, pic = ps
        if text_of(label).strip() != "Completed Results:" or label.find(".//" + qn("w:drawing")) is not None \
                or pic.find(".//" + qn("w:drawing")) is None or text_of(pic).strip():
            continue
        first = pic.find(qn("w:r"))
        for k, (tag, val) in enumerate((("t", "Completed Results:"), ("br", None))):
            r = pic.makeelement(qn("w:r"), {})
            rpr = r.makeelement(qn("w:rPr"), {})
            rpr.append(rpr.makeelement(qn("w:rStyle"), {qn("w:val"): "Heading4Char"}))
            r.append(rpr)
            el = r.makeelement(qn("w:" + tag), {})
            if val:
                el.text = val
            r.append(el)
            first.addprevious(r)
        tc.remove(label)
        n += 1
    return n


def part_by_hash(d, sha1_prefix):
    """An embedded image part by the start of its content sha1 (stable across builds, unlike media names)."""
    import hashlib
    hits = [p for p in d.part.package.parts if "/media/" in str(p.partname)
            and hashlib.sha1(p.blob).hexdigest().startswith(sha1_prefix)]
    assert len(hits) == 1, (sha1_prefix, len(hits))
    return hits[0]


def renumber_callout(part, bbox, n, fill=None):
    """Repaint one orange numbered circle (bbox in image px) with a new number, in place. The circle is
    redrawn supersampled so its edge matches; the old shadow stays."""
    import io
    from PIL import ImageDraw, ImageFont
    im = Image.open(io.BytesIO(part.blob)).convert("RGB")
    x0, y0, x1, y1 = bbox
    fill = fill or im.getpixel(((x0 + x1) // 2, y0 + 3))
    m, k = 3, 4
    patch = im.crop((x0 - m, y0 - m, x1 + m + 1, y1 + m + 1))
    big = patch.resize((patch.width * k, patch.height * k), Image.LANCZOS)
    dr = ImageDraw.Draw(big)
    dr.ellipse((m * k, m * k, (x1 - x0 + m + 1) * k, (y1 - y0 + m + 1) * k), fill=fill)
    font = next((ImageFont.truetype(f, int((y1 - y0) * 0.72 * k)) for f in ("segoeuib.ttf", "arialbd.ttf")), None)
    cx, cy = (x1 - x0 + 1) / 2 + m, (y1 - y0 + 1) / 2 + m
    dr.text((cx * k, (cy + 0.5) * k), str(n), font=font, fill="white", anchor="mm")
    im.paste(big.resize(patch.size, Image.LANCZOS), (x0 - m, y0 - m))
    buf = io.BytesIO()
    im.save(buf, "PNG", dpi=im.info.get("dpi", (96, 96)))
    part._blob = buf.getvalue()


def set_row_picture_width(d, heading, row_starts, width_in):
    """Set the width (keeping aspect) of the picture in the instruction row whose text starts with
    row_starts: for David's own 96-dpi images placed in a rebuilt table (fit_images leaves them alone)."""
    _, tbl_el, _ = find_table(d, heading)
    tr = next(r for r in tbl_el.findall(qn("w:tr")) if text_of(r).startswith(row_starts))
    drawing = next(tr.iter(qn("w:drawing")))
    ext = next(drawing.iter(_WP + "extent"))
    cx = int(width_in * EMU)
    _set_size(drawing, cx, int(cx * int(ext.get("cy")) / int(ext.get("cx"))))


def replace_row_picture(d, heading, row_starts, png, width_in=None):
    """Put a new picture in one instruction row only (the old one may be shared with other rows, as
    image45 is between Connect to Data and Athena). Replaces the picture in the row's last cell, or adds
    one if that cell has none. fit_images sizes it afterwards if it is a 192-dpi capture."""
    _, tbl_el, t = find_table(d, heading)
    tr = next(r for r in tbl_el.findall(qn("w:tr")) if text_of(r).startswith(row_starts))
    tc = tr.findall(qn("w:tc"))[-1]
    for r in [r for r in tc.iter(qn("w:r")) if r.find(".//" + qn("w:drawing")) is not None]:
        r.getparent().remove(r)
    cell = docx.table._Cell(tc, t)
    im = Image.open(png)
    w = width_in or min(im.width / im.info.get("dpi", (96, 96))[0], CELL_MAX)
    cell.paragraphs[0].add_run().add_picture(str(png), width=Inches(w))
