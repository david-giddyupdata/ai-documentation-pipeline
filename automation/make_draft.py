"""Build `QuickSight Journey 1 - draft.docx` from the master plus every finished exercise.

The master is only read. Re-run after editing any ex_*.py or re-annotating screenshots.
"""
import json
import shutil
from pathlib import Path

import docx

import ex_add_area
import ex_add_clustered
import ex_add_combo
import ex_add_donut
import ex_add_filled
import ex_add_heat
import ex_add_line
import ex_add_pivot
import ex_add_point
import ex_add_scatter
import ex_add_sidebyside
import ex_add_spark
import ex_add_stacked
import ex_add_table
import ex_add_tree
import ex_add_word
import ex_background
import ex_connect_data
import ex_create_analysis
import ex_athena
import ex_front
import ex_s3
import ex_redshift
import ex_postgres
import ex_rva
import ex_row_level_s4
import ex_icon
import ex_lacs
import ex_row_level
import ex_table_calcs
import ex_themes

from build_lib import (fix_numbering, fit_images, keep_picture_rows_together, merge_completed_results,
                       page_breaks_to_headings,
                       space_after_pictures)

HERE = Path(__file__).resolve().parent
FIT_REPORT = HERE / "_work" / "fit_report.json"
MASTER = HERE.parent / "QuickSight Journey 1.docx"
DRAFT = HERE.parent / "QuickSight Journey 1 - draft.docx"
EXERCISES = [ex_row_level, ex_connect_data, ex_create_analysis, ex_athena, ex_front, ex_s3, ex_redshift,
             ex_postgres, ex_rva, ex_row_level_s4, ex_table_calcs, ex_lacs, ex_background, ex_icon, ex_themes,
             ex_add_table, ex_add_pivot, ex_add_heat,
             ex_add_clustered, ex_add_stacked, ex_add_sidebyside,
             ex_add_line, ex_add_spark, ex_add_area, ex_add_combo,
             ex_add_donut, ex_add_tree, ex_add_word,
             ex_add_filled, ex_add_point, ex_add_scatter]

if __name__ == "__main__":
    shutil.copy2(MASTER, DRAFT)
    d = docx.Document(DRAFT)
    for ex in EXERCISES:
        ex.apply(d)
        print("applied", ex.HEADING)
    fix_numbering(d)  # after the exercises, so it also covers steps they rebuilt
    print("fixed step numbering")
    # layout passes (David's review rules): picture size, space below pictures, no blank pages,
    # rows with a picture kept on one page
    print("Completed Results rows merged:", merge_completed_results(d))
    short = fit_images(d)
    FIT_REPORT.write_text(json.dumps(short, indent=1, ensure_ascii=False), encoding="utf-8")
    print(f"sized pictures; {len(short)} can't reach 90% (see {FIT_REPORT.name})")
    print("space below pictures:", space_after_pictures(d))
    print("page breaks moved onto headings:", page_breaks_to_headings(d))
    print("rows kept on one page:", keep_picture_rows_together(d))
    d.save(DRAFT)
    print("saved", DRAFT)
